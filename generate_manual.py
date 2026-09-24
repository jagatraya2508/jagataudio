# -*- coding: utf-8 -*-
"""Generate Jagat Audio tutorial/manual as Word (.docx) and PDF."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent
VERSION = "2.1.2"
DOCX_PATH = ROOT / f"Tutorial_Manual_JagatAudio_v{VERSION}.docx"
PDF_PATH = ROOT / f"Tutorial_Manual_JagatAudio_v{VERSION}.pdf"
SCREENSHOT = Path(
    r"C:\Users\wisnu\.cursor\projects\d-Programer-jagataudio\assets"
    r"\c__Users_wisnu_AppData_Roaming_Cursor_User_workspaceStorage"
    r"_c8240561f800334d3a2b30c9f32fa6f4_images_image-b98da3de-f173-4d93-93e1-b3ff94befa81.png"
)

PURPLE = RGBColor(0x83, 0x38, 0xEC)
PINK = RGBColor(0xFF, 0x47, 0x7E)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x4B, 0x55, 0x63)

FONT_DIR = Path(r"C:\Windows\Fonts")
FONT_REG = FONT_DIR / "arial.ttf"
FONT_BOLD = FONT_DIR / "arialbd.ttf"
FONT_ITALIC = FONT_DIR / "ariali.ttf"
FONT_BI = FONT_DIR / "arialbi.ttf"


def set_run_font(run, name="Calibri", size=11, bold=False, color=DARK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D1D5DB")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = PURPLE if level == 1 else PINK
        run.font.name = "Calibri"
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text)
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(item, style=style)
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run_font(run, size=11)


def add_callout(doc, title, text, fill="EEF2FF"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=PURPLE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(text)
    set_run_font(r2, size=10, color=GRAY)
    doc.add_paragraph()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "8338EC")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            if ri % 2 == 1:
                shade_cell(cell, "F5F3FF")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=10)
    doc.add_paragraph()


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("JAGAT AUDIO")
    set_run_font(r, size=32, bold=True, color=PURPLE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI Stem Separation & Karaoke")
    set_run_font(r, size=16, color=PINK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Tutorial & Manual Pengguna  •  Versi {VERSION}")
    set_run_font(r, size=13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Panduan lengkap menu, fungsi, dan cara memakai setiap fitur")
    set_run_font(r, size=11, color=GRAY)

    if SCREENSHOT.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(SCREENSHOT), width=Inches(6.3))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dokumen ini mengikuti label menu di aplikasi: Stem Separator, Web Audio Converter, DAW Studio, Guitar Gear Detector, serta Media Playlist, Lirik & Cord.")
    set_run_font(r, size=10, color=GRAY)

    doc.add_page_break()

    add_heading_styled(doc, "Daftar Isi", 1)
    toc = [
        "1. Tentang Jagat Audio",
        "2. Instalasi & Persiapan",
        "3. Login, Registrasi, dan Lisensi",
        "4. Peta Menu Utama",
        "5. Stem Separator (Pemisahan AI & Karaoke)",
        "6. Web Audio Converter",
        "7. DAW Studio",
        "8. Guitar Gear Detector",
        "9. Media Playlist, Lirik & Cord",
        "10. Cari Tab Online",
        "11. Tips, Batasan, dan Troubleshooting",
        "12. Ringkasan Alur Kerja Cepat",
    ]
    add_bullets(doc, toc, numbered=True)

    # 1
    add_heading_styled(doc, "1. Tentang Jagat Audio", 1)
    add_body(
        doc,
        "Jagat Audio adalah aplikasi desktop all-in-one untuk musisi, penyanyi karaoke, gitaris, dan produser pemula. "
        "Aplikasi memisahkan lagu menjadi beberapa instrumen (stem) dengan AI lokal (Demucs HT 6-stems), "
        "memutar dan mencampur hasilnya, merekam vokal karaoke, mengunduh audio/video dari tautan web, "
        "merekam dan mixing multi-track di DAW Studio, mendeteksi perkiraan gear gitar dari rekaman lagu, "
        "serta memutar playlist sambil menampilkan lirik dan chord.",
    )
    add_body(doc, "Jagat Audio memproses pemisahan stem di komputer Anda (bukan di cloud). Kecepatan bergantung pada spesifikasi PC. Fitur pencarian lirik, chord, unduhan web, dan deteksi gear membutuhkan koneksi internet.")

    add_heading_styled(doc, "Apa yang bisa Anda lakukan", 2)
    add_bullets(doc, [
        "Karaoke: mute stem Vokal, ubah nada/tempo, nyanyi pakai mic, rekam, lalu export.",
        "Latihan band: isolasi drum, bass, gitar, atau piano.",
        "Mixing cepat: volume, pan, Auto Balance, Audio Enhancer, Vocal Processing.",
        "Produksi: rekaman multi-track, drum/bass generator, efek, mastering, export WAV/MP3.",
        "Latihan gitar: chord/tab, transposisi sesuai nada dasar, deteksi gear.",
        "Pustaka media: playlist folder/file, lirik LRC, video karaoke.",
    ])

    # 2
    add_heading_styled(doc, "2. Instalasi & Persiapan", 1)
    add_heading_styled(doc, "Yang Anda terima dari admin", 2)
    add_bullets(doc, [
        "Folder aplikasi (biasanya JagatAudio.exe beserta file pendukung).",
        "File lisensi berakhiran .lic yang terikat ke Hardware ID komputer Anda.",
    ])
    add_heading_styled(doc, "Cara menjalankan", 2)
    add_bullets(doc, [
        "Extract arsip ZIP jika ada, lalu taruh folder di lokasi tetap (misalnya D:\\JagatAudio).",
        "Jangan pindah-pindah folder aplikasi setelah lisensi aktif kecuali Anda tahu konsekuensinya.",
        "Double-click JagatAudio.exe. Aplikasi akan membuka antarmuka (browser internal/lokal).",
        "Tunggu beberapa detik saat pertama kali karena model AI dan layanan lokal sedang dimuat.",
    ], numbered=True)
    add_callout(
        doc,
        "Catatan lisensi",
        "File .lic bersifat unik per komputer. Tidak bisa disalin ke PC lain. Jika ganti hardware signifikan, minta lisensi baru dengan Hardware ID terbaru.",
        "FFF3CD",
    )

    # 3
    add_heading_styled(doc, "3. Login, Registrasi, dan Lisensi", 1)
    add_heading_styled(doc, "3.1 Aktivasi lisensi (layar kunci)", 2)
    add_body(doc, "Jika lisensi belum valid, layar menampilkan Aktivasi Lisensi Diperlukan.")
    add_bullets(doc, [
        "Salin Hardware ID Anda (klik teks ID). Status berubah menjadi Tersalin ke clipboard.",
        "Kirim Hardware ID ke admin/penjual Jagat Audio.",
        "Terima file .lic, lalu seret ke area Upload File Lisensi (.lic) atau klik untuk memilih file.",
        "Tunggu Mengaktifkan lisensi.... Jika sukses, aplikasi membuka menu utama.",
    ], numbered=True)

    add_heading_styled(doc, "3.2 Tombol Info", 2)
    add_body(doc, "Di bawah judul Jagat Audio terdapat tombol Info (ikon kunci). Klik untuk membuka kartu lisensi:")
    add_bullets(doc, [
        "Badge Lisensi Aktif, jenis lisensi, sisa hari, dan tanggal kedaluwarsa (Exp).",
        "Salin Hardware ID untuk perpanjangan.",
        "Perpanjang / ganti file lisensi (.lic) tanpa keluar aplikasi.",
        "Sembunyikan untuk merapikan tampilan.",
    ])

    add_heading_styled(doc, "3.3 Login dan daftar akun", 2)
    add_body(doc, "Setelah lisensi aktif, Anda masuk ke akun lokal aplikasi.")
    add_bullets(doc, [
        "Masuk ke Akun Anda: isi Username dan Password, lalu login.",
        "Belum punya akun? klik Daftar Sekarang. Isi Username, Email, dan Password, lalu Daftar.",
        "Sudah punya akun? klik Login di Sini.",
    ])
    add_body(doc, "Akun dipakai untuk sesi fitur (pemisahan, unduhan, pencarian). Jika muncul pesan sesi habis, login ulang.")

    add_heading_styled(doc, "3.4 Manajemen User (khusus Admin)", 2)
    add_body(doc, "Akun admin dapat membuka Manajemen User: Tambah User, ubah username/role, atau hapus user. Pengguna biasa tidak perlu menu ini.")

    # 4
    add_heading_styled(doc, "4. Peta Menu Utama", 1)
    add_body(doc, "Lima tab di bilah navigasi (urutan sama seperti di aplikasi):")
    add_table(
        doc,
        ["Menu", "Kegunaan singkat"],
        [
            ["Stem Separator", "Unggah lagu/video, pisahkan 6 stem AI, mixer, karaoke, export."],
            ["Web Audio Converter", "Cari atau paste tautan, preview, unduh Audio MP3 atau Video Klip MP4."],
            ["DAW Studio", "Studio rekaman multi-track, mixing, mastering, drum & bass generator."],
            ["Guitar Gear Detector", "Analisis file audio untuk tebakan tone, gitar, amp, dan pedal."],
            ["Media Playlist, Lirik & Cord", "Putar folder/file, lirik sinkron, chord, transposisi, video karaoke."],
        ],
    )
    add_body(doc, "Saat mixer Stem atau DAW tampil penuh layar, tombol Menu Utama (panah kiri) mengembalikan Anda ke beranda tab tersebut.")

    # 5
    add_heading_styled(doc, "5. Stem Separator (Pemisahan AI & Karaoke)", 1)
    add_body(doc, "Tab default. Memisahkan audio menjadi 6 channel: Vokal, Drum, Bass, Gitar, Piano, dan Lainnya.")

    add_heading_styled(doc, "5.1 Beranda: Upload Baru vs Proyek Tersimpan", 2)
    add_body(doc, "Ada dua sub-tab:")
    add_bullets(doc, [
        "Upload Baru — pilih file baru untuk dipisah.",
        "Proyek Tersimpan — buka hasil pemisahan sebelumnya tanpa memproses ulang. Anda bisa ubah nama (pensil) atau hapus (tong sampah).",
    ])

    add_heading_styled(doc, "5.2 Unggah dan proses", 2)
    add_bullets(doc, [
        "Klik Pilih File. Format didukung: MP3, WAV, dan video (MP4, MOV, AVI, MKV, WEBM).",
        "Pratinjau lagu/video, isi Nama Proyek (contoh: Bon Jovi - Its My Life). Nama tersimpan ke Proyek Tersimpan.",
        "Opsional: Cari Tab Online (oranye) untuk chord/tab sebelum atau tanpa menunggu pemisahan selesai.",
        "Klik Mulai Pemisahan. Status: Mengunggah Lagu → Memisahkan Audio (Demucs HT 6-Stems) → analisis nada dasar → Memuat Hasil.",
        "Pilih Ulang untuk ganti file. Jika gagal, Coba Lagi.",
    ], numbered=True)
    add_callout(
        doc,
        "Sabar saat memproses",
        "Progres menampilkan persen dan sisa waktu. Di 100% aplikasi masih bisa menganalisis nada dasar dari stem bass, piano, dan gitar. Jangan tutup aplikasi sampai mixer muncul.",
        "EEF2FF",
    )

    add_heading_styled(doc, "5.3 Mixer studio", 2)
    add_body(doc, "Setelah siap, tampilan penuh Mixer Instrumen. Jika sumbernya video, video tampil (suara dari stem, video di-mute).")
    add_heading_styled(doc, "Kontrol master", 3)
    add_bullets(doc, [
        "Play / Pause dan timeline seek.",
        "Export Media (audio) atau Export MP4 (jika sumber video).",
        "Potong Lagu / Trim ON: atur Mulai dan Akhir, lalu Simpan Potongan.",
        "Nada Dasar: pilih kunci; lagu ikut pindah. Keterangan asli ... muncul jika pitch ≠ 0.",
        "Pitch: −12 sampai +12 semitone.",
        "Tempo: 50% sampai 150%.",
    ])

    add_heading_styled(doc, "Rekaman Karaoke", 3)
    add_bullets(doc, [
        "Aktifkan Mic (izinkan akses mikrofon). Gunakan headset agar tidak feedback.",
        "Mute channel Vokal di mixer, atur Mic Vol.",
        "Sync delay (0–300 ms): naikkan jika vokal hasil export tertinggal dari backing.",
        "Mulai Rekam / Stop. Indikator REC menampilkan durasi.",
        "Pitch Coach: jarum Rendah / Pas / Tinggi membandingkan nada mic dengan lagu.",
        "Voice Effect: Reverb, Echo, Chorus, Pitch, Robot (plus preset seperti Natural, Hall, Radio, Chipmunk, Deep, Thick).",
        "Unduh WebM untuk menyimpan rekaman.",
    ])

    add_heading_styled(doc, "Mixer Instrumen", 3)
    add_bullets(doc, [
        "Enam fader dB (−60 sampai +12). Nilai dB bisa diketik lalu Enter.",
        "Pan L–C–R per channel.",
        "Mute (ikon speaker) per instrumen — mute Vokal = backing karaoke.",
        "Auto Balance: mengatur volume & pan agar mix lebih seimbang secara otomatis.",
    ])

    add_heading_styled(doc, "Audio Enhancer", 3)
    add_body(doc, "Panel mastering cepat untuk mix keseluruhan:")
    add_bullets(doc, [
        "EQ cepat: Bass, Mid, Treble (−12 sampai +12 dB) dan Master Vol.",
        "Compressor, Limiter, Normalize, Denoise (on/off).",
        "Reverb, Delay, Warmth, Widener.",
        "Graphic EQ 10 pita: 31 Hz sampai 16 kHz.",
    ])

    add_heading_styled(doc, "Vocal Processing (khusus stem Vokal)", 3)
    add_bullets(doc, [
        "Auto Magic: preset pemrosesan vokal industri.",
        "Noise Gate: meredam derau saat tidak ada vokal.",
        "Leveler + Target: meratakan dinamika volume vokal.",
    ])

    add_heading_styled(doc, "Lirik Karaoke (panel kanan)", 3)
    add_body(doc, "Saat lagu diputar, lirik diupayakan dimuat otomatis. Lirik bertanda LRC punya timestamp (karaoke). Anda dapat mencari ulang jika judul salah, dan menyesuaikan sinkron (termasuk double-klik baris yang sedang dinyanyikan). Tempo mixer sudah diperhitungkan.")

    # 6
    add_heading_styled(doc, "6. Web Audio Converter", 1)
    add_body(doc, "Mengunduh audio atau video klip dari tautan web (YouTube dan sumber lain yang didukung, misalnya SoundCloud pada hasil pencarian).")
    add_bullets(doc, [
        "Pilih jenis: Audio MP3 atau Video Klip MP4.",
        "Paste tautan, atau ketik judul (contoh: Coldplay Yellow). Tombol mikrofon mendukung pencarian suara (Chrome).",
        "Klik Cari Lagu atau Cari Video.",
        "Hasil pencarian: preview Play/Pause, seek, mundur/maju 10 detik, kembali ke awal.",
        "Unduh MP3 / Unduh MP4. Tunggu persiapan dan progres mengunduh.",
        "Jika sukses: Simpan MP3/MP4, atau Unduh Lainnya. Jika gagal: Coba Lagi / Batal.",
    ], numbered=True)
    add_callout(
        doc,
        "Hak cipta",
        "Gunakan unduhan hanya untuk keperluan pribadi yang sah. Hormati hak cipta pemilik konten.",
        "FFF3CD",
    )

    # 7
    add_heading_styled(doc, "7. DAW Studio", 1)
    add_body(doc, "Studio rekaman/arrangement terinspirasi alur kerja DAW modern. Dari beranda DAW:")
    add_bullets(doc, [
        "Project Baru — timeline kosong (default dua track).",
        "Import Audio — pilih satu atau banyak file MP3/WAV/OGG/FLAC/AAC/M4A, langsung masuk project baru.",
        "Project Tersimpan — buka project sebelumnya (BPM, jumlah track, tanggal). Ikon tong sampah untuk hapus.",
        "Kembali ke Menu — keluar ke tab aplikasi.",
    ])

    add_heading_styled(doc, "7.1 Transport bar", 2)
    add_table(
        doc,
        ["Kontrol", "Fungsi"],
        [
            ["Menu Utama", "Kembali ke menu Jagat Audio."],
            ["Ikon folder", "DAW Home / buka project."],
            ["Stop / Play / Record", "Stop (Enter), Play/Pause (Space), Record (R)."],
            ["Loop & Metronome", "Mengulang wilayah; ketukan BPM."],
            ["Ikon drum", "Generate Drum Loop & Simulator."],
            ["Tanda ~", "Generate Bassline (nada bisa terisi dari Stem Separator)."],
            ["BPM & time signature", "Tempo 20–300; birama misalnya 4/4."],
            ["Output perangkat", "Pilih speaker/headphone. Hindari output Communications. Ada panduan Valeton GP-200."],
            ["Nama project", "Klik untuk ganti nama."],
            ["Simpan / Open / Export", "Simpan (Ctrl+S), buka, export WAV atau MP3."],
        ],
    )

    add_heading_styled(doc, "7.2 Toolbar edit", 2)
    add_bullets(doc, [
        "Pointer (V): pilih dan geser region.",
        "Split (C): potong clip di posisi kursor.",
        "Eraser (E): hapus clip.",
        "Undo / Redo (Ctrl+Z / Ctrl+Shift+Z).",
        "Snap: Bar, 1/2, 1/4, 1/8, 1/16, atau Off.",
        "Zoom horizontal dan Height track.",
    ])

    add_heading_styled(doc, "7.3 Track dan mixer", 2)
    add_bullets(doc, [
        "Nama track, urutkan naik/turun, hapus track. Klik kanan header untuk menu konteks.",
        "Arm record, mute, solo, monitor (Monitor DAW ON bisa echo jika amp/interface juga bunyi — matikan monitor jika pakai GP-200 + headphone unit).",
        "Volume, pan, dan slot efek per track.",
        "Efek termasuk Guitar Amp Cabinet Simulator (Clean Tube / Crunch Overdrive / Heavy Lead), compressor, filter, delay, hall reverb, pitch shift (semitone).",
        "Mixer channel MST / MAIN OUT plus MASTER BUS MASTERING SUITE: glue compressor, linear EQ, stereo imager & mono check, limiter/output ceiling.",
        "Preset mixing profesional tersedia di sistem (misalnya Acoustic Guitar Shimmer) melalui panel mixing.",
    ])

    add_heading_styled(doc, "7.4 Drum Loop & Bassline", 2)
    add_body(doc, "Generate Drum Loop: pilih Drum Kit (mis. Standard Acoustic), pola, fill (Tanpa Fill, Classic Tom Fill, Trap Hat Roll), dan format output. Generate Bassline: tangga mayor/minor, gaya (mis. Funky Groove), mengikuti nada dasar yang disarankan dari Stem Separator jika ada.")

    add_callout(
        doc,
        "Rekaman gitar + GP-200",
        "Iringan boleh keluar ke GP-200. Rekaman hanya mengambil input gitar. Monitor DAW OFF agar tidak echo; dengar gitar dari headphone unit.",
        "EEF2FF",
    )

    # 8
    add_heading_styled(doc, "8. Guitar Gear Detector", 1)
    add_body(doc, "AI memperkirakan profil tone gitar dari file audio (bukan jaminan merk pasti, melainkan referensi mixing/tone-matching).")
    add_bullets(doc, [
        "Sub-tab Analisis Baru atau Riwayat Analisis.",
        "Opsional: isi Nama Artis / Judul Lagu untuk membantu akurasi.",
        "Pilih File Audio, pratinjau di Preview Lagu Asli.",
        "Klik Mulai Analisis Tone.",
        "Hasil: Akurasi/Confidence (%), Tipe Tone, Detail Gitar (model, pickup, senar), Amplifier (head, cab, setting), Efek & Pedal, deskripsi, dan Alternatif Gear Lain.",
        "Riwayat: buka ulang, ganti nama, atau hapus analisis lama.",
    ], numbered=True)

    # 9
    add_heading_styled(doc, "9. Media Playlist, Lirik & Cord", 1)
    add_body(doc, "Pemutar pustaka media dengan lirik dan chord. Chord di UI ditulis Cord sesuai label menu aplikasi.")

    add_heading_styled(doc, "9.1 Player vs Kelola Playlist", 2)
    add_body(doc, "Beranda player: putar lagu, lirik, chord. Tombol Kelola Playlist atau Playlist Baru membuka halaman manajemen.")

    add_heading_styled(doc, "Dua cara memilih media", 3)
    add_bullets(doc, [
        "Pilih 1 Folder Penuh: memindai folder + subfolder. Dialog Windows bisa menyembunyikan daftar file — itu normal; klik Select Folder.",
        "Pilih File Media (1 / Acak / Rentang): 1 file, Ctrl untuk beberapa file acak, Shift untuk rentang.",
    ])
    add_body(doc, "Format: MP3, MP4, M4A, WAV, OGG, FLAC, AAC, WEBM. Tip: pilih folder lewat Chrome/Edge agar lirik bisa ditulis langsung ke folder MP3 (indikator Folder siap menyimpan lirik).")

    add_heading_styled(doc, "Playlist Saya", 3)
    add_bullets(doc, [
        "Playlist Baru → nama → Buat.",
        "Terapkan: memutar grup tanpa wajib simpan.",
        "Simpan / Update: agar tetap ada setelah aplikasi ditutup (alur mirip playlist tetap).",
        "Panah: lihat daftar lagu; Pindah ke... memindahkan lagu ke grup lain.",
        "Kosongkan Player menghapus sesi putar saat ini.",
        "Kembali ke Player untuk memutar.",
    ])

    add_heading_styled(doc, "9.2 Pemutar", 2)
    add_bullets(doc, [
        "Shuffle, Previous, Play/Pause, Next, Repeat (mati / ulang semua / ulang 1 lagu).",
        "Seekbar waktu. Video MP4 tampil di area player.",
        "Nada Dasar: deteksi kunci; pilih kunci baru untuk transposisi. Label perkiraan jika deteksi tidak pasti.",
        "Daftar trek: klik untuk putar, ikon lirik/chord jika sudah ada, hapus dari playlist.",
        "Cari lagu, chord, atau lirik di kotak atas (contoh: Metallica Nothing Else Matters) — file MP3 tidak wajib (mode songbook).",
    ])

    add_heading_styled(doc, "9.3 Panel Lirik", 2)
    add_bullets(doc, [
        "Tab Lirik: karaoke word-by-word jika file LRC; teks biasa jika tanpa timestamp.",
        "Cari Lirik Manual jika judul/penyanyi salah.",
        "Cari versi ber-timestamp (.lrc) jika hanya teks polos.",
        "Anda dapat menaruh file NamaLagu.lrc di folder yang sama.",
        "Tangga nada: tombol naik/turun semitone (−12…+12) dan reset.",
    ])

    add_heading_styled(doc, "9.4 Panel Chord", 2)
    add_bullets(doc, [
        "Tab Chord memuat chord/tab (bisa otomatis dari judul lagu).",
        "Chord ikut transposisi saat nada dasar/pitch diubah.",
        "Unduh Chord (.txt) untuk disimpan.",
    ])

    add_heading_styled(doc, "9.5 Sinkronkan Video Karaoke", 2)
    add_body(doc, "Jika ada lirik ber-timestamp, Anda dapat menyesuaikan offset (detik) lalu mengekspor/membakar lirik ke video karaoke. Geser trim jika teks terlalu maju atau mundur terhadap vokal.")

    # 10
    add_heading_styled(doc, "10. Cari Tab Online", 1)
    add_body(doc, "Tersedia dari Stem Separator (tombol oranye Cari Tab Online) dan dari pencarian chord di playlist.")
    add_bullets(doc, [
        "Isi atau perbaiki judul lagu dan artis.",
        "Klik Cari Sekarang.",
        "Baca tabulatur di layar; unduh sebagai .txt bila tersedia.",
        "Butuh internet. Hasil bergantung pada ketersediaan di sumber online.",
    ], numbered=True)

    # 11
    add_heading_styled(doc, "11. Tips, Batasan, dan Troubleshooting", 1)
    add_table(
        doc,
        ["Masalah", "Yang bisa dicoba"],
        [
            ["Layar Aktivasi Lisensi", "Pastikan .lic untuk Hardware ID ini. Salin ID, minta file baru ke admin."],
            ["Login gagal / sesi habis", "Periksa username/password. Login ulang."],
            ["Pemisahan lama / hang", "Tunggu ETA. Jangan sleep PC. Tutup aplikasi berat lain. RAM/CPU minim memperlambat Demucs."],
            ["Gagal memproses / unduh", "Cek internet, Coba Lagi, ganti file/tautan."],
            ["Mic tidak bunyi / feedback", "Izinkan mic di OS. Pakai headset. Turunkan Mic Vol. Mute stem Vokal."],
            ["Karaoke tidak sinkron", "Naikkan Sync delay. Di playlist, geser offset/trim LRC."],
            ["Pitch aneh / chipmunk", "Kembalikan Pitch ke 0 atau Nada Dasar asli. Tempo 100%."],
            ["Lirik/chord salah lagu", "Cari Lirik Manual / ketik artis+judul yang benar."],
            ["Folder playlist tidak terlihat filenya", "Normal di dialog Windows; tetap Select Folder."],
            ["Echo di DAW", "Matikan Monitor DAW; jangan pakai output Communications."],
            ["Lisensi hampir habis", "Info → sisa hari oranye. Upload .lic perpanjangan."],
        ],
    )
    add_heading_styled(doc, "Praktik terbaik", 2)
    add_bullets(doc, [
        "Beri Nama Proyek yang rapi agar mudah dicari di Proyek Tersimpan.",
        "Simpan playlist yang sering dipakai (jangan hanya Terapkan).",
        "Export mix setelah mute/volume/pitch/tempo sudah pas.",
        "Backup file .lic di tempat aman (tetap hanya untuk PC yang sama).",
    ])

    # 12
    add_heading_styled(doc, "12. Ringkasan Alur Kerja Cepat", 1)
    add_heading_styled(doc, "Karaoke 10 menit", 2)
    add_bullets(doc, [
        "Stem Separator → Upload → Mulai Pemisahan.",
        "Mute Vokal → set Nada Dasar/Pitch jika perlu.",
        "Aktifkan Mic → headset → Mulai Rekam → Stop → Export.",
    ], numbered=True)
    add_heading_styled(doc, "Latihan gitar sambil lihat chord", 2)
    add_bullets(doc, [
        "Media Playlist → Pilih File/Folder atau Cari judul.",
        "Tab Chord; sesuaikan Nada Dasar.",
        "Atau Stem Separator → Cari Tab Online, lalu mute gitar di mixer untuk play-along.",
    ], numbered=True)
    add_heading_styled(doc, "Cover + rekam gitar", 2)
    add_bullets(doc, [
        "Pisahkan lagu di Stem Separator, catat nada dasar.",
        "DAW Studio → Import Audio (stem) atau rekam baru.",
        "Generate drum/bass bila perlu, mixing, Export MP3/WAV.",
    ], numbered=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"\n© 2026 Jagat Audio  •  Manual pengguna v{VERSION}")
    set_run_font(r, size=10, color=GRAY)

    doc.save(str(DOCX_PATH))
    return DOCX_PATH


class ManualPDF(FPDF):
    def header(self):
        if self.page_no() == 1:
            return
        self.set_font("Arial", "B", 9)
        self.set_text_color(131, 56, 236)
        self.cell(0, 8, "Jagat Audio  |  Tutorial & Manual Pengguna", align="L")
        self.set_text_color(157, 78, 221)
        self.cell(0, 8, f"v{VERSION}", align="R", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(131, 56, 236)
        self.line(18, 16, 192, 16)
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.set_text_color(107, 114, 128)
        self.cell(0, 10, f"Halaman {self.page_no()}  |  Jagat Audio {VERSION}", align="C")


def pdf_fonts(pdf):
    if FONT_REG.exists():
        pdf.add_font("Arial", "", str(FONT_REG))
        pdf.add_font("Arial", "B", str(FONT_BOLD if FONT_BOLD.exists() else FONT_REG))
        pdf.add_font("Arial", "I", str(FONT_ITALIC if FONT_ITALIC.exists() else FONT_REG))
        pdf.add_font("Arial", "BI", str(FONT_BI if FONT_BI.exists() else FONT_REG))


def h1(pdf, text):
    pdf.ln(3)
    pdf.set_font("Arial", "B", 16)
    pdf.set_text_color(131, 56, 236)
    pdf.multi_cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(255, 71, 126)
    pdf.line(pdf.l_margin, pdf.get_y(), 192, pdf.get_y())
    pdf.ln(3)


def h2(pdf, text):
    pdf.ln(2)
    pdf.set_font("Arial", "B", 13)
    pdf.set_text_color(255, 71, 126)
    pdf.multi_cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)


def h3(pdf, text):
    pdf.set_font("Arial", "B", 11)
    pdf.set_text_color(75, 85, 99)
    pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)


def body(pdf, text):
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Arial", "", 10)
    pdf.set_text_color(26, 26, 46)
    pdf.multi_cell(0, 5.4, text, new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin)
    pdf.ln(1.5)


def bullets(pdf, items, numbered=False):
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Arial", "", 10)
    pdf.set_text_color(26, 26, 46)
    for i, item in enumerate(items, 1):
        prefix = f"{i}. " if numbered else "- "
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 5.2, prefix + item, new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin)
    pdf.ln(1.5)


def callout(pdf, title, text):
    pdf.set_fill_color(238, 242, 255)
    pdf.set_font("Arial", "B", 10)
    pdf.set_text_color(131, 56, 236)
    pdf.multi_cell(0, 6, title, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 9)
    pdf.set_text_color(75, 85, 99)
    pdf.multi_cell(0, 5, text, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)


def simple_table(pdf, headers, rows):
    col_w = [50, 124]
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    col_w[1] = usable - col_w[0]
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Arial", "B", 9)
    pdf.set_fill_color(131, 56, 236)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(col_w[0], 7, headers[0], border=1, fill=True)
    pdf.cell(col_w[1], 7, headers[1], border=1, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 8.5)
    pdf.set_text_color(26, 26, 46)
    for ri, row in enumerate(rows):
        pdf.set_fill_color(245, 243, 255) if ri % 2 else pdf.set_fill_color(255, 255, 255)
        pdf.set_x(pdf.l_margin)
        y0 = pdf.get_y()
        h0 = pdf.multi_cell(col_w[0], 5, row[0], dry_run=True, output="HEIGHT")
        h1 = pdf.multi_cell(col_w[1], 5, row[1], dry_run=True, output="HEIGHT")
        h = max(h0, h1) + 1
        if y0 + h > pdf.h - 20:
            pdf.add_page()
            y0 = pdf.get_y()
        pdf.set_xy(pdf.l_margin, y0)
        pdf.rect(pdf.l_margin, y0, col_w[0], h, style="DF")
        pdf.rect(pdf.l_margin + col_w[0], y0, col_w[1], h, style="DF")
        pdf.set_xy(pdf.l_margin + 1, y0 + 0.5)
        pdf.multi_cell(col_w[0] - 2, 5, row[0], new_x="RIGHT", new_y="TOP")
        pdf.set_xy(pdf.l_margin + col_w[0] + 1, y0 + 0.5)
        pdf.multi_cell(col_w[1] - 2, 5, row[1], new_x="LMARGIN", new_y="NEXT")
        pdf.set_xy(pdf.l_margin, y0 + h)
    pdf.set_x(pdf.l_margin)
    pdf.ln(3)


def build_pdf():
    pdf = ManualPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)
    pdf_fonts(pdf)
    pdf.add_page()

    pdf.ln(18)
    pdf.set_font("Arial", "B", 28)
    pdf.set_text_color(131, 56, 236)
    pdf.cell(0, 12, "JAGAT AUDIO", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "B", 14)
    pdf.set_text_color(255, 71, 126)
    pdf.cell(0, 8, "AI Stem Separation & Karaoke", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "B", 12)
    pdf.set_text_color(26, 26, 46)
    pdf.cell(0, 8, f"Tutorial & Manual Pengguna  •  Versi {VERSION}", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 10)
    pdf.set_text_color(75, 85, 99)
    pdf.multi_cell(0, 6, "Panduan lengkap menu, fungsi, dan cara memakai setiap fitur sesuai tampilan aplikasi.", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    if SCREENSHOT.exists():
        pdf.image(str(SCREENSHOT), x=18, w=174)
        pdf.set_x(pdf.l_margin)
    pdf.ln(6)
    pdf.set_font("Arial", "I", 9)
    pdf.multi_cell(
        0,
        5,
        "Dokumen ini mengikuti label menu: Stem Separator, Web Audio Converter, DAW Studio, Guitar Gear Detector, dan Media Playlist, Lirik & Cord.",
        align="C",
        new_x="LMARGIN",
        new_y="NEXT",
    )

    pdf.add_page()
    h1(pdf, "Daftar Isi")
    bullets(pdf, [
        "1. Tentang Jagat Audio",
        "2. Instalasi & Persiapan",
        "3. Login, Registrasi, dan Lisensi",
        "4. Peta Menu Utama",
        "5. Stem Separator (Pemisahan AI & Karaoke)",
        "6. Web Audio Converter",
        "7. DAW Studio",
        "8. Guitar Gear Detector",
        "9. Media Playlist, Lirik & Cord",
        "10. Cari Tab Online",
        "11. Tips, Batasan, dan Troubleshooting",
        "12. Ringkasan Alur Kerja Cepat",
    ], numbered=False)

    h1(pdf, "1. Tentang Jagat Audio")
    body(
        pdf,
        "Jagat Audio adalah aplikasi desktop all-in-one untuk musisi, penyanyi karaoke, gitaris, dan produser pemula. "
        "Aplikasi memisahkan lagu menjadi beberapa instrumen (stem) dengan AI lokal (Demucs HT 6-stems), "
        "memutar dan mencampur hasilnya, merekam vokal karaoke, mengunduh audio/video dari tautan web, "
        "merekam dan mixing multi-track di DAW Studio, mendeteksi perkiraan gear gitar dari rekaman lagu, "
        "serta memutar playlist sambil menampilkan lirik dan chord.",
    )
    body(pdf, "Pemisahan stem berjalan di komputer Anda. Kecepatan bergantung pada spesifikasi PC. Fitur pencarian lirik, chord, unduhan web, dan deteksi gear membutuhkan internet.")
    h2(pdf, "Apa yang bisa Anda lakukan")
    bullets(pdf, [
        "Karaoke: mute stem Vokal, ubah nada/tempo, nyanyi pakai mic, rekam, lalu export.",
        "Latihan band: isolasi drum, bass, gitar, atau piano.",
        "Mixing cepat: volume, pan, Auto Balance, Audio Enhancer, Vocal Processing.",
        "Produksi: rekaman multi-track, drum/bass generator, efek, mastering, export WAV/MP3.",
        "Latihan gitar: chord/tab, transposisi, deteksi gear.",
        "Pustaka media: playlist, lirik LRC, video karaoke.",
    ])

    h1(pdf, "2. Instalasi & Persiapan")
    h2(pdf, "Yang Anda terima dari admin")
    bullets(pdf, [
        "Folder aplikasi (JagatAudio.exe beserta file pendukung).",
        "File lisensi .lic yang terikat ke Hardware ID komputer Anda.",
    ])
    h2(pdf, "Cara menjalankan")
    bullets(pdf, [
        "Extract ZIP bila ada, taruh folder di lokasi tetap (contoh D:\\JagatAudio).",
        "Double-click JagatAudio.exe dan tunggu layanan lokal dimuat.",
        "Jangan memindahkan folder aplikasi setelah lisensi aktif tanpa keperluan.",
    ], numbered=True)
    callout(pdf, "Catatan lisensi", "File .lic unik per komputer. Ganti hardware signifikan berarti minta lisensi baru dengan Hardware ID terbaru.")

    h1(pdf, "3. Login, Registrasi, dan Lisensi")
    h2(pdf, "3.1 Aktivasi lisensi")
    body(pdf, "Jika lisensi belum valid, layar menampilkan Aktivasi Lisensi Diperlukan.")
    bullets(pdf, [
        "Salin Hardware ID (klik teks ID).",
        "Kirim ke admin/penjual, terima file .lic.",
        "Seret ke Upload File Lisensi (.lic) atau klik untuk memilih.",
        "Tunggu aktivasi; menu utama terbuka jika sukses.",
    ], numbered=True)
    h2(pdf, "3.2 Tombol Info")
    bullets(pdf, [
        "Di bawah judul Jagat Audio: Info membuka kartu Lisensi Aktif, jenis, sisa hari, Exp.",
        "Salin Hardware ID; Perpanjang / ganti file lisensi (.lic); Sembunyikan untuk merapikan.",
    ])
    h2(pdf, "3.3 Login dan daftar")
    bullets(pdf, [
        "Masuk ke Akun Anda: Username + Password.",
        "Daftar Sekarang: Username, Email, Password.",
        "Jika sesi habis, login ulang.",
    ])
    h2(pdf, "3.4 Manajemen User")
    body(pdf, "Hanya akun Admin: Tambah User, ubah role, hapus user.")

    h1(pdf, "4. Peta Menu Utama")
    simple_table(pdf, ["Menu", "Kegunaan singkat"], [
        ["Stem Separator", "Unggah lagu/video, pisahkan 6 stem AI, mixer, karaoke, export."],
        ["Web Audio Converter", "Cari/paste tautan, preview, unduh MP3 atau MP4."],
        ["DAW Studio", "Rekaman multi-track, mixing, mastering, drum & bass generator."],
        ["Guitar Gear Detector", "Tebakan tone, gitar, amp, dan pedal dari file audio."],
        ["Media Playlist, Lirik & Cord", "Putar folder/file, lirik, chord, transposisi, video karaoke."],
    ])
    body(pdf, "Saat mixer Stem atau DAW penuh layar, tombol Menu Utama mengembalikan ke beranda tab.")

    h1(pdf, "5. Stem Separator")
    body(pdf, "Memisahkan audio menjadi Vokal, Drum, Bass, Gitar, Piano, dan Lainnya.")
    h2(pdf, "5.1 Upload Baru dan Proyek Tersimpan")
    bullets(pdf, [
        "Upload Baru: pilih file baru.",
        "Proyek Tersimpan: buka hasil lama tanpa proses ulang; pensil = ganti nama, tong sampah = hapus.",
    ])
    h2(pdf, "5.2 Unggah dan proses")
    bullets(pdf, [
        "Pilih File (MP3, WAV, MP4, MOV, AVI, MKV, WEBM).",
        "Pratinjau, isi Nama Proyek, opsional Cari Tab Online.",
        "Mulai Pemisahan. Tunggu unggah, Demucs 6-stems, analisis nada dasar, lalu mixer.",
        "Pilih Ulang atau Coba Lagi jika perlu.",
    ], numbered=True)
    callout(pdf, "Sabar saat memproses", "Di 100% aplikasi masih bisa menganalisis nada dasar. Jangan tutup sampai mixer muncul.")
    h2(pdf, "5.3 Mixer studio")
    h3(pdf, "Kontrol master")
    bullets(pdf, [
        "Play/Pause, timeline, Export Media atau Export MP4.",
        "Potong Lagu (Trim): Mulai–Akhir lalu Simpan Potongan.",
        "Nada Dasar, Pitch −12…+12, Tempo 50%–150%.",
    ])
    h3(pdf, "Rekaman Karaoke")
    bullets(pdf, [
        "Aktifkan Mic, headset, mute Vokal, atur Mic Vol dan Sync delay (0–300 ms).",
        "Mulai Rekam / Stop, Pitch Coach (Rendah/Pas/Tinggi), Voice Effect, Unduh WebM.",
    ])
    h3(pdf, "Mixer Instrumen")
    bullets(pdf, [
        "Fader dB, pan L–C–R, mute per stem, Auto Balance.",
    ])
    h3(pdf, "Audio Enhancer & Vocal Processing")
    bullets(pdf, [
        "Bass/Mid/Treble, Master Vol, Compressor, Limiter, Normalize, Denoise, Reverb, Delay, Warmth, Widener, EQ 10 pita.",
        "Vocal Processing: Auto Magic, Noise Gate, Leveler + Target (khusus stem Vokal).",
    ])
    h3(pdf, "Lirik Karaoke")
    body(pdf, "Lirik dimuat otomatis bila memungkinkan. Badge LRC = ada timestamp. Cari ulang jika judul salah. Double-klik baris untuk bantu sinkron.")

    h1(pdf, "6. Web Audio Converter")
    bullets(pdf, [
        "Pilih Audio MP3 atau Video Klip MP4.",
        "Paste tautan atau ketik judul; mikrofon = pencarian suara (Chrome).",
        "Cari Lagu/Video, preview, lalu Unduh MP3/MP4.",
        "Simpan file, atau Unduh Lainnya / Coba Lagi.",
    ], numbered=True)
    callout(pdf, "Hak cipta", "Gunakan unduhan hanya untuk keperluan pribadi yang sah.")

    h1(pdf, "7. DAW Studio")
    bullets(pdf, [
        "Project Baru, Import Audio, Project Tersimpan, Kembali ke Menu.",
        "Transport: Stop, Play (Space), Record (R), Loop, Metronome, Drum Loop, Bassline (~), BPM, output perangkat, Simpan, Open, Export WAV/MP3.",
        "Tool: Pointer (V), Split (C), Eraser (E), Undo/Redo, Snap, Zoom, Height.",
        "Track: nama, urutan, arm, mute, solo, monitor, volume, pan, FX (amp sim, compressor, delay, reverb, pitch).",
        "Master: MAIN OUT, glue compressor, EQ, stereo imager, limiter.",
        "Monitor DAW OFF jika pakai GP-200 agar tidak echo. Hindari output Communications.",
    ])

    h1(pdf, "8. Guitar Gear Detector")
    bullets(pdf, [
        "Analisis Baru: opsional isi artis/judul, Pilih File Audio, preview, Mulai Analisis Tone.",
        "Hasil: confidence, tipe tone, gitar, amp, pedal, deskripsi, alternatif.",
        "Riwayat Analisis: buka, ganti nama, hapus.",
        "Ini perkiraan referensi tone, bukan jaminan merek pasti.",
    ], numbered=True)

    h1(pdf, "9. Media Playlist, Lirik & Cord")
    body(pdf, "Pemutar pustaka + lirik + chord. Label menu memakai ejaan Cord.")
    h2(pdf, "Kelola Playlist")
    bullets(pdf, [
        "Pilih 1 Folder Penuh (subfolder ikut; dialog Windows boleh kosong — tetap Select Folder).",
        "Pilih File Media: 1 file, Ctrl acak, Shift rentang.",
        "Playlist Baru, Terapkan (tanpa simpan), Simpan/Update, Pindah ke..., Kosongkan Player.",
        "Chrome/Edge memudahkan menulis lirik ke folder MP3.",
    ])
    h2(pdf, "Player")
    bullets(pdf, [
        "Shuffle, prev/next, Repeat all/one/off, seek, video MP4.",
        "Nada Dasar + tangga nada −12…+12.",
        "Cari judul tanpa file MP3 (songbook).",
        "Tab Lirik (LRC karaoke / teks) dan Chord (transposisi, Unduh Chord .txt).",
        "Sinkronkan Video Karaoke dengan offset/trim detik.",
    ])

    h1(pdf, "10. Cari Tab Online")
    bullets(pdf, [
        "Dari Stem Separator (tombol oranye) atau pencarian chord playlist.",
        "Isi judul & artis, Cari Sekarang, baca/unduh .txt. Butuh internet.",
    ], numbered=True)

    h1(pdf, "11. Tips dan Troubleshooting")
    simple_table(pdf, ["Masalah", "Yang bisa dicoba"], [
        ["Aktivasi Lisensi", "Pastikan .lic sesuai Hardware ID ini; minta file baru ke admin."],
        ["Login / sesi habis", "Cek password, login ulang."],
        ["Pemisahan lama", "Tunggu ETA, jangan sleep PC, tutup aplikasi berat."],
        ["Gagal proses/unduh", "Cek internet, Coba Lagi, ganti file/tautan."],
        ["Mic / feedback", "Headset, izinkan mic OS, turunkan Mic Vol, mute Vokal."],
        ["Karaoke tidak sinkron", "Naikkan Sync delay atau offset LRC."],
        ["Pitch aneh", "Pitch 0, Nada Dasar asli, Tempo 100%."],
        ["Lirik/chord salah", "Cari Lirik Manual / ketik artis+judul benar."],
        ["Echo DAW", "Monitor OFF; hindari output Communications."],
        ["Lisensi hampir habis", "Info → upload .lic perpanjangan."],
    ])

    h1(pdf, "12. Ringkasan Alur Kerja Cepat")
    h2(pdf, "Karaoke")
    bullets(pdf, ["Stem Separator → Upload → Mulai Pemisahan → mute Vokal → mic → rekam → Export."], numbered=True)
    h2(pdf, "Latihan gitar")
    bullets(pdf, ["Playlist → Chord / Nada Dasar, atau Stem Separator → Cari Tab Online lalu mute gitar."], numbered=True)
    h2(pdf, "Cover + rekam")
    bullets(pdf, ["Pisahkan stem → DAW Import/rekam → drum/bass generator → mixing → Export MP3/WAV."], numbered=True)

    pdf.set_font("Arial", "I", 9)
    pdf.set_text_color(107, 114, 128)
    pdf.ln(8)
    pdf.cell(0, 6, f"© 2026 Jagat Audio  •  Manual pengguna v{VERSION}", align="C")

    pdf.output(str(PDF_PATH))
    return PDF_PATH


if __name__ == "__main__":
    d = build_docx()
    p = build_pdf()
    print(f"Wrote {d}")
    print(f"Wrote {p}")
