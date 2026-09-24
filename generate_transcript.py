import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        
def create_transcript():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Transkrip Video TikTok: Review Aplikasi Jagat Audio (Lengkap)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("Tema: Aplikasi AI multifungsi untuk memisahkan vokal, membuat musik (DAW), mencari chord, & converter.")
    doc.add_paragraph("Durasi Estimasi: ~60-90 detik\n")
    
    # Hook
    add_heading(doc, '1. Hook / Intro (0:00 - 0:05)', 2)
    doc.add_paragraph('Visual: Tampilkan ekspresi kaget atau antusias, lalu tunjukkan layar PC dengan tampilan awal Jagat Audio.')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Suka nyanyi karaoke, anak band yang seneng ngulik lagu, atau bahkan produser musik pemula? Kalian wajib banget cobain aplikasi super komplit ini!"')
    
    # Menu 1
    add_heading(doc, '2. Menu Login & Aktivasi Lisensi (0:05 - 0:12)', 2)
    doc.add_paragraph('Visual: Rekam layar (screen record) saat proses login, lalu arahkan kursor ke menu aktivasi lisensi.')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Kenalin, ini Jagat Audio! Tinggal bikin akun, login, dan pastiin lisensinya aktif ya, biar bisa bebas eksplor semua fitur gokilnya."')

    # Menu 2
    add_heading(doc, '3. Menu Utama: AI Stem Separator & Mixer (0:12 - 0:28)', 2)
    doc.add_paragraph('Visual: Sorot tab "Stem Separator", drag & drop MP3, lalu ke layar "Mixer Player". Tunjukkan saat menekan tombol Mute di vokal, dan mengatur pitch (nada) serta tempo.')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Fitur pertamanya ada AI Stem Separator! Tinggal masukin lagu, dan BOOM... kepisah otomatis jadi 6 instrumen. Kalian bisa mute vokal buat karaoke, ubah nada, sampai ubah tempo lagunya sebelum di-ekspor. Canggih kan?"')

    # Menu 3
    add_heading(doc, '4. Fitur DAW Studio (Digital Audio Workstation) (0:28 - 0:42)', 2)
    doc.add_paragraph('Visual: Beralih ke layar DAW Studio (tampilan timeline multi-track, drum loop generator, dan mixing efek).')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Buat kalian yang mau bikin musik sendiri, di sini juga udah ada fitur DAW Studio-nya lho! Kalian bisa rekaman multi-track, mixing pakai efek, dan bahkan ada fitur Drum Generator instan. Nggak perlu lagi software berat!"')

    # Menu 4
    add_heading(doc, '5. Web Audio Converter (0:42 - 0:50)', 2)
    doc.add_paragraph('Visual: Sorot menu "Web Audio Converter". Tunjukkan proses memilih file dan convert format audio secara kilat.')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Kalau butuh ubah format file audio biar kompatibel, ada Web Audio Converter bawaan. Proses convert-nya cepet banget langsung di dalam aplikasinya, nggak butuh pindah-pindah web."')

    # Menu 5
    add_heading(doc, '6. Fitur Cari Tab Online & YouTube to MP3 (0:50 - 1:05)', 2)
    doc.add_paragraph('Visual: Tunjukkan sekilas pencarian "Cari Tab Online", dilanjutkan tab "YouTube to MP3".')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Nggak cuma itu, kalian juga bisa cari chord dan tab gitar otomatis dari lagu yang diputar, plus ada fitur YouTube to MP3 buat ekstrak atau misahin vokal langsung dari link YouTube. Praktis abis!"')

    # Outro
    add_heading(doc, '7. Outro / Call to Action (1:05 - 1:15)', 2)
    doc.add_paragraph('Visual: Menampilkan kamu tersenyum sambil menunjuk ke arah bio/link atau layar.')
    p = doc.add_paragraph()
    p.add_run('Audio/Voiceover: ').bold = True
    p.add_run('"Bener-bener all-in-one banget buat musisi dan penyuka musik! Yuk, buruan cobain Jagat Audio sekarang! Jangan lupa like, save, dan share video ini ke temen ngeband kamu!"')

    doc.save('Transkrip_Tiktok_JagatAudio_V2.docx')
    print("Document saved as Transkrip_Tiktok_JagatAudio_V2.docx")

if __name__ == '__main__':
    create_transcript()
